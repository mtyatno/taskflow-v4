"""Ekspor DOCX hanya me-resolve drawing milik pengekspor (spec §3, baris 3e).

`_make_drawing_resolver(uid)` dipakai ekspor note tersimpan (GET /api/scratchpad/{id}/export/docx) dan
ekspor live (POST /api/scratchpad/export/docx, isinya bebas dari klien). Token direktif ::draw[...]
di-resolve lewat id angka, client_id, lalu judul (persis, lalu heuristik "gambar/canvas -"), dan semuanya
terscope ke user_id pengekspor: drawing user lain (walau di-pin), judul milik user lain, dan fallback lama
"id drawing = id note" tidak pernah ter-resolve. Token berbentuk id angka atau UUID yang tidak ter-resolve
tidak ditebak lewat judul (akan menyematkan drawing pengekspor sendiri yang kebetulan berjudul sama).

DB tes dipakai bersama satu sesi pytest: username, judul, client_id, dan marker dibuat unik.
"""
import io
import re
import uuid
from types import SimpleNamespace

import docx

import webapp
from conftest import db, register_user

UNKNOWN = "00000000-0000-0000-0000-000000000000"


def _uniq():
    """Tag unik TANPA angka (huruf a-p). Heuristik judul resolver (dipertahankan) mengubah judul bawaan
    "Gambar <token>" jadi LIKE '%token%' atas judul drawing milik pengekspor: token id angka ("Gambar 2")
    akan kebetulan cocok dengan judul drawing A yang memuat "2" bila tag-nya memuat angka (tes jadi flaky)."""
    return uuid.uuid4().hex[:10].translate(str.maketrans("0123456789", "ghijklmnop"))


def _auth(client, username):
    """Daftarkan user dan kembalikan header Bearer-nya. Cookie sesi hasil register dibuang
    supaya request berikutnya hanya membawa header (aman dipakai bergantian antar user)."""
    data = register_user(client, username, f"{username}@test.id")
    client.cookies.clear()
    return {"Authorization": f"Bearer {data.get('token') or data.get('access_token')}"}


def _uid(username):
    """id user dari tabel users."""
    conn = db()
    try:
        return conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()["id"]
    finally:
        conn.close()


def _svg(marker):
    return f"<svg><text>{marker}</text></svg>"


def _drawing(client, headers, title, svg, cid):
    """Buat drawing lewat API; kembalikan baris JSON-nya (id, client_id, title, svg_preview, ...)."""
    r = client.post("/api/drawings", headers=headers, json={
        "title": title, "data_json": '{"shapes":{}}', "svg_preview": svg, "client_id": cid,
    })
    assert r.status_code == 200, r.text
    return r.json()


def _scenario(client):
    """User A dan B, masing-masing satu drawing (svg marker, client_id, judul unik). Drawing B di-pin."""
    tag = _uniq()
    ha = _auth(client, f"expusera_{tag}")
    hb = _auth(client, f"expuserb_{tag}")
    a = _drawing(client, ha, f"denah-a-{tag}", _svg(f"MARKER-A-{tag}"), f"cid-a-{tag}")
    b = _drawing(client, hb, f"denah-b-{tag}", _svg(f"MARKER-B-{tag}"), f"cid-b-{tag}")
    pin = client.patch(f"/api/drawings/{b['id']}/pin", headers=hb)
    assert pin.status_code == 200 and pin.json()["is_pinned"] == 1, pin.text
    return SimpleNamespace(
        tag=tag, ha=ha, uid_a=_uid(f"expusera_{tag}"), uid_b=_uid(f"expuserb_{tag}"), a=a, b=b,
        svg_a=_svg(f"MARKER-A-{tag}"),
    )


def _create_note(client, headers, title, content, list_id=None):
    """Buat note (pribadi, atau di workspace bila list_id diisi); kembalikan id-nya."""
    payload = {"title": title, "content": content}
    if list_id is not None:
        payload["list_id"] = list_id
    r = client.post("/api/scratchpad", json=payload, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _workspace(client, headers, name):
    """Buat workspace (shared list) milik user; kembalikan id-nya."""
    r = client.post("/api/lists", json={"name": name}, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _add_member(list_id, user_id):
    """Jadikan user anggota workspace langsung lewat DB (alur undangan tidak relevan untuk tes ini)."""
    conn = db()
    try:
        conn.execute(
            "INSERT INTO list_members (list_id, user_id, joined_at) VALUES (?, ?, '2026-10-18T00:00:00')",
            (list_id, user_id),
        )
        conn.commit()
    finally:
        conn.close()


def _plant_drawing_with_next_note_id(client, headers, owner_id, title, svg):
    """Sisipkan drawing milik owner_id dengan id yang akan dipakai note BERIKUTNYA (autoincrement
    scratchpad_notes), supaya fallback lama "id drawing = id note" punya sasaran. Kembalikan id itu."""
    _create_note(client, headers, f"seed {title}", "x")  # pastikan baris sqlite_sequence scratchpad_notes ada
    conn = db()
    try:
        did = max(
            conn.execute("SELECT COALESCE(MAX(id), 0) FROM drawings").fetchone()[0],
            conn.execute("SELECT COALESCE(MAX(id), 0) FROM scratchpad_notes").fetchone()[0],
            conn.execute("SELECT COALESCE(MAX(seq), 0) FROM sqlite_sequence WHERE name = 'scratchpad_notes'").fetchone()[0],
        ) + 1000
        conn.execute(
            "INSERT INTO drawings (id, user_id, title, data_json, svg_preview, is_pinned, created_at, updated_at) "
            "VALUES (?, ?, ?, '{}', ?, 0, '2026-10-18T00:00:00', '2026-10-18T00:00:00')",
            (did, owner_id, title, svg),
        )
        conn.execute("UPDATE sqlite_sequence SET seq = ? WHERE name = 'scratchpad_notes'", (did - 1,))
        conn.commit()
        return did
    finally:
        conn.close()


def _exports(client, headers, note_id, content):
    """Ekspor content lewat kedua endpoint DOCX (note tersimpan dan live); kembalikan
    {endpoint: (teks semua paragraf, jumlah gambar di dokumen)}. Drawing yang ter-resolve jadi gambar
    dengan keterangan berisi judulnya; yang tidak ter-resolve jadi paragraf "[Gambar/Canvas: ...]"."""
    saved = client.get(f"/api/scratchpad/{note_id}/export/docx", headers=headers)
    live = client.post("/api/scratchpad/export/docx", headers=headers, json={
        "note_id": note_id, "title": "Ekspor live", "content": content,
    })
    views = {}
    for kind, r in (("tersimpan", saved), ("live", live)):
        assert r.status_code == 200, (kind, r.text)
        doc = docx.Document(io.BytesIO(r.content))
        views[kind] = ("\n".join(p.text for p in doc.paragraphs), len(doc.inline_shapes))
    return views


def _placeholders(text):
    """Judul pada paragraf "[Gambar/Canvas: judul]": tepat rujukan drawing yang tidak ter-resolve."""
    return re.findall(r"\[Gambar/Canvas: ([^\]]*)\]", text)


# ── Resolver ─────────────────────────────────────────────────────────────────

def test_export_resolver_rejects_other_users_drawings(client):
    """Drawing user lain (walau di-pin) tidak ter-resolve lewat id, client_id, judul direktif,
    heuristik "gambar/canvas -", maupun token berisi judul."""
    s = _scenario(client)
    res = webapp._make_drawing_resolver(s.uid_a)

    attempts = {
        "id": res(str(s.b["id"])),
        "client_id": res(s.b["client_id"]),
        "judul direktif": res("tidak-ada", title=s.b["title"]),
        "judul dengan prefix gambar": res("tidak-ada", title=f"Gambar - {s.b['title']}"),
        "token judul": res(s.b["title"]),
    }
    leaked = [name for name, row in attempts.items() if row is not None]
    assert not leaked, f"drawing user lain ter-resolve lewat: {leaked}"


def test_export_resolver_resolves_own_drawings(client):
    """Drawing milik pengekspor ter-resolve lewat id, client_id (dukungan baru), judul direktif persis,
    dan heuristik lama "gambar/canvas -", dengan bentuk hasil {id, title, svg} yang dibaca docx_exporter."""
    s = _scenario(client)
    res = webapp._make_drawing_resolver(s.uid_a)
    expected = {"id": s.a["id"], "title": s.a["title"], "svg": s.svg_a}

    got = {
        "id": res(str(s.a["id"])),
        "client_id": res(s.a["client_id"]),
        "judul direktif": res("tidak-ada", title=s.a["title"]),
        "judul dengan prefix gambar": res("tidak-ada", title=f"Gambar - {s.a['title']}"),
    }
    assert got == {name: expected for name in got}


def test_export_resolver_no_note_id_fallback(client):
    """Token tak dikenal tidak jatuh ke drawing milik user lain yang id-nya kebetulan sama dengan id note
    yang diekspor (fallback lama "id = note_id"). Dicek di resolver dan lewat kedua endpoint ekspor,
    karena fallback itu hanya aktif bila pemanggil menyerahkan note_id, seperti endpoint dulu."""
    s = _scenario(client)
    planted_id = _plant_drawing_with_next_note_id(
        client, s.ha, s.uid_b, f"denah-b-nota-{s.tag}", _svg(f"MARKER-B-NOTE-{s.tag}"))
    content = f"::draw[{UNKNOWN}]"
    note_id = _create_note(client, s.ha, f"Note A {s.tag}", content)
    assert note_id == planted_id, "skenario salah: id note A harus sama dengan id drawing B"

    assert webapp._make_drawing_resolver(s.uid_a)(UNKNOWN) is None

    for kind, (text, pictures) in _exports(client, s.ha, note_id, content).items():
        assert pictures == 0, f"{kind}: drawing B (id sama dengan id note) ikut tersemat"
        assert _placeholders(text) == [f"Gambar {UNKNOWN}"], kind


# ── Endpoint ekspor ──────────────────────────────────────────────────────────

def test_export_endpoints_embed_only_exporters_drawings(client):
    """Kedua endpoint ekspor DOCX hanya menyematkan drawing milik pengekspor (lewat id, client_id,
    judul direktif); rujukan ke drawing user lain, walau di-pin, tampil sebagai placeholder."""
    s = _scenario(client)
    content = "\n\n".join([
        f"::draw[{s.a['id']}]",
        f"::draw[{s.a['client_id']}]",
        f'::draw[tidak-ada]{{title="{s.a["title"]}"}}',
        f"::draw[{s.b['id']}]",
        f"::draw[{s.b['client_id']}]",
        f'::draw[tidak-ada]{{title="{s.b["title"]}"}}',
    ])
    note_id = _create_note(client, s.ha, f"Note A {s.tag}", content)
    unresolved = sorted([f"Gambar {s.b['id']}", f"Gambar {s.b['client_id']}", s.b["title"]])

    for kind, (text, pictures) in _exports(client, s.ha, note_id, content).items():
        assert pictures == 3, kind
        assert text.count(s.a["title"]) == 3, kind  # keterangan ketiga gambar itu: drawing A
        assert sorted(_placeholders(text)) == unresolved, kind  # tepat tiga rujukan ke drawing B yang tidak ter-resolve


# ── Token id/UUID tidak pernah ditebak lewat judul ───────────────────────────

def test_export_resolver_does_not_guess_by_title_for_id_tokens(client):
    """Aplikasi menulis direktif ::draw[<id atau uuid>]{title="<judul drawing>"} (judul bawaan "Untitled Drawing").
    Token id/UUID yang tidak ter-resolve untuk pengekspor (drawing milik orang lain, mis. note bersama pemilik lain)
    tidak boleh jatuh ke drawing PENGEKSPOR SENDIRI yang kebetulan berjudul sama. Token biasa tetap memakai judul."""
    tag = _uniq()
    ha = _auth(client, f"expguessa_{tag}")
    hm = _auth(client, f"expguessm_{tag}")
    uid_m = _uid(f"expguessm_{tag}")
    a = _drawing(client, ha, "Untitled Drawing", _svg(f"MARKER-GA-{tag}"), str(uuid.uuid4()))
    m = _drawing(client, hm, "Untitled Drawing", _svg(f"MARKER-GM-{tag}"), str(uuid.uuid4()))
    res = webapp._make_drawing_resolver(uid_m)

    assert res(str(a["id"]), title="Untitled Drawing") is None
    assert res(a["client_id"], title="Untitled Drawing") is None
    assert res(a["client_id"].upper(), title="Untitled Drawing") is None
    # kontrol: drawing sendiri tetap ter-resolve lewat id, dan token biasa tetap memakai heuristik judul
    assert res(str(m["id"]), title="Untitled Drawing")["id"] == m["id"]
    assert res("tidak-ada", title="Untitled Drawing")["id"] == m["id"]


def test_export_member_of_shared_note_does_not_get_own_lookalike_drawing(client):
    """M anggota workspace A. Note bersama A merujuk drawing pribadi A lewat id dan UUID dengan judul bawaan
    "Untitled Drawing" (persis seperti yang ditulis aplikasi). Ekspor DOCX oleh M tidak boleh menyematkan drawing M
    sendiri yang kebetulan berjudul sama: tidak ada gambar, hanya placeholder. Kontrol: rujukan ke drawing M sendiri
    di note M tetap menjadi satu gambar."""
    tag = _uniq()
    ha = _auth(client, f"expmema_{tag}")
    m_name = f"expmemm_{tag}"
    hm = _auth(client, m_name)
    a = _drawing(client, ha, "Untitled Drawing", _svg(f"MARKER-MA-{tag}"), str(uuid.uuid4()))
    m = _drawing(client, hm, "Untitled Drawing", _svg(f"MARKER-MM-{tag}"), str(uuid.uuid4()))
    w = _workspace(client, ha, f"Workspace {tag}")
    _add_member(w, _uid(m_name))
    content = "\n\n".join([
        f'::draw[{a["id"]}]{{title="Untitled Drawing"}}',
        f'::draw[{a["client_id"]}]{{title="Untitled Drawing"}}',
    ])
    shared_note = _create_note(client, ha, f"Note bersama A {tag}", content, list_id=w)

    for kind, (text, pictures) in _exports(client, hm, shared_note, content).items():
        assert pictures == 0, f"{kind}: drawing M sendiri tersemat untuk rujukan ke drawing A"
        assert _placeholders(text) == ["Untitled Drawing"] * 2, kind

    own_content = f'::draw[{m["id"]}]{{title="Untitled Drawing"}}'
    own_note = _create_note(client, hm, f"Note M {tag}", own_content)
    for kind, (text, pictures) in _exports(client, hm, own_note, own_content).items():
        assert pictures == 1, kind
